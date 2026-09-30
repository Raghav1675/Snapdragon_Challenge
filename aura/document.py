from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

import fitz
import pandas as pd
from PIL import Image

try:
    import pytesseract
except Exception:  # pragma: no cover
    pytesseract = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None


@dataclass
class DocumentContent:
    name: str
    source_type: str
    text: str
    pages: list[str]
    metadata: dict
    original_bytes: bytes | None = None


def normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def _extract_pdf(data: bytes) -> DocumentContent:
    doc = fitz.open(stream=data, filetype="pdf")
    pages: list[str] = []
    for page in doc:
        pages.append(normalize_text(page.get_text("text")))
    text = "\n\n".join(p for p in pages if p)
    meta = dict(doc.metadata or {})
    meta["page_count"] = len(pages)
    return DocumentContent("uploaded.pdf", "PDF", text, pages, meta, data)


def _extract_docx(data: bytes) -> DocumentContent:
    if DocxDocument is None:
        raise RuntimeError("python-docx is not installed")
    doc = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            tables.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = normalize_text("\n".join(paragraphs + tables))
    return DocumentContent("uploaded.docx", "DOCX", text, [text], {}, data)


def _extract_image(data: bytes, name: str) -> DocumentContent:
    image = Image.open(io.BytesIO(data))
    text = ""
    if pytesseract is not None:
        try:
            text = pytesseract.image_to_string(image)
        except Exception:
            text = ""
    text = normalize_text(text)
    meta = {"image_size": image.size, "ocr_available": bool(text)}
    return DocumentContent(name, "IMAGE", text, [text], meta, data)


def _extract_csv(data: bytes, name: str) -> DocumentContent:
    df = pd.read_csv(io.BytesIO(data))
    text = df.to_csv(index=False)
    return DocumentContent(name, "CSV", normalize_text(text), [normalize_text(text)], {"rows": len(df), "columns": len(df.columns)}, data)


def extract_document(data: bytes, filename: str) -> DocumentContent:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        result = _extract_pdf(data)
    elif suffix == ".docx":
        result = _extract_docx(data)
    elif suffix in {".txt", ".md", ".markdown"}:
        text = normalize_text(data.decode("utf-8", errors="replace"))
        result = DocumentContent(filename, "TEXT", text, [text], {}, data)
    elif suffix in {".csv"}:
        result = _extract_csv(data, filename)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}:
        result = _extract_image(data, filename)
    else:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")
    result.name = filename
    return result


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 160) -> list[str]:
    text = normalize_text(text)
    if not text:
        return []
    words = text.split()
    chunks: list[str] = []
    step = max(1, chunk_size - overlap)
    for start in range(0, len(words), step):
        chunk = " ".join(words[start : start + chunk_size])
        if chunk:
            chunks.append(chunk)
        if start + chunk_size >= len(words):
            break
    return chunks


def redact_pdf_bytes(data: bytes) -> bytes:
    """Redact detected sensitive strings from searchable PDF text and return a new PDF."""
    from .privacy import scan_text

    source = fitz.open(stream=data, filetype="pdf")
    for page in source:
        page_text = page.get_text("text")
        findings = scan_text(page_text)
        for finding in findings:
            try:
                rects = page.search_for(finding.value)
            except Exception:
                rects = []
            for rect in rects:
                page.add_redact_annot(rect, fill=(0, 0, 0))
        if findings:
            page.apply_redactions()
    output = io.BytesIO()
    source.save(output)
    source.close()
    return output.getvalue()
